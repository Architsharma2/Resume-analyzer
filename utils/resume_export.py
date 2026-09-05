from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def colors_for(style: str):
    if style == "classic":
        return RGBColor(0, 0, 0), RGBColor(40, 40, 40), RGBColor(90, 90, 90), "000000"
    if style == "minimal":
        return RGBColor(45, 55, 72), RGBColor(40, 40, 40), RGBColor(100, 100, 100), "718096"
    return RGBColor(31, 78, 121), RGBColor(35, 35, 35), RGBColor(90, 90, 90), "1F4E79"


def set_run(run, size=10.5, bold=False, color=RGBColor(35, 35, 35), name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_bottom_line(paragraph, color="1F4E79"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_heading(doc, text, style):
    primary, _, _, line = colors_for(style)
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    set_run(run, 11, True, primary)
    add_bottom_line(p, line)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def body_para(doc, text, size=10.5, bold=False, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, size, bold, color or RGBColor(35, 35, 35))
    p.paragraph_format.space_after = Pt(2)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run(run, 10.5, False, RGBColor(40, 40, 40))
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.15)


def title_date_line(doc, left_main, left_sub, right_date, style):
    _, dark, gray, _ = colors_for(style)
    p = doc.add_paragraph()
    r1 = p.add_run(left_main)
    set_run(r1, 10.5, True, dark)
    if left_sub:
        r2 = p.add_run(f"  ·  {left_sub}")
        set_run(r2, 10.5, False, gray)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r3 = p.add_run(f"\t{right_date}" if right_date else "")
    set_run(r3, 10, False, gray)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(4)


def parse_block_lines(text):
    items = []
    for line in (text or "").split("\n"):
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in line.split("|")]
        while len(parts) < 4:
            parts.append("")
        title, org, dates, bullets = parts[:4]
        bullet_list = [b.strip() for b in bullets.split(";") if b.strip()]
        items.append({"title": title, "org": org, "dates": dates, "bullets": bullet_list})
    return items


def export_docx(data, photo_bytes=None) -> bytes:
    style = (data.get("style") or "modern").lower()
    primary, dark, gray, _ = colors_for(style)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    if photo_bytes:
        table = doc.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells

        p_photo = left.paragraphs[0]
        run = p_photo.add_run()
        run.add_picture(BytesIO(photo_bytes), width=Inches(1.05))

        p_name = right.paragraphs[0]
        r = p_name.add_run((data.get("name") or "YOUR NAME").upper())
        set_run(r, 18, True, primary)

        p_role = right.add_paragraph()
        rr = p_role.add_run(data.get("headline") or data.get("role") or "")
        set_run(rr, 10.5, False, dark)

        contacts = [x for x in [
            data.get("location"), data.get("phone"), data.get("email"),
            data.get("linkedin"), data.get("github"), data.get("portfolio")
        ] if x]
        p_c = right.add_paragraph()
        rc = p_c.add_run("  ·  ".join(contacts))
        set_run(rc, 9.5, False, gray)
    else:
        p_name = doc.add_paragraph()
        r = p_name.add_run((data.get("name") or "YOUR NAME").upper())
        set_run(r, 20, True, primary)
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER

        p_role = doc.add_paragraph()
        rr = p_role.add_run(data.get("headline") or data.get("role") or "")
        set_run(rr, 10.5, False, dark)
        p_role.alignment = WD_ALIGN_PARAGRAPH.CENTER

        contacts = [x for x in [
            data.get("location"), data.get("phone"), data.get("email"),
            data.get("linkedin"), data.get("github"), data.get("portfolio")
        ] if x]
        p_c = doc.add_paragraph()
        rc = p_c.add_run("  ·  ".join(contacts))
        set_run(rc, 9.5, False, gray)
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    objective = (data.get("summary") or "").strip()
    if not objective:
        objective = (
            f"Motivated student targeting {data.get('role') or 'entry-level'} opportunities. "
            f"Eager to learn, contribute, and grow through real projects and teamwork."
        )
    section_heading(doc, "Objective", style)
    body_para(doc, objective)

    section_heading(doc, "Skills", style)
    for label, key in [
        ("Technical", "skills"),
        ("Tools", "tools"),
        ("Soft Skills", "soft_skills"),
        ("Languages", "languages"),
    ]:
        val = (data.get(key) or "").strip()
        if not val:
            continue
        p = doc.add_paragraph()
        r1 = p.add_run(f"{label}: ")
        set_run(r1, 10.5, True, primary)
        r2 = p.add_run(val)
        set_run(r2, 10.5, False, dark)

    section_heading(doc, "Experience", style)
    exp_items = parse_block_lines(data.get("experience", ""))
    if exp_items:
        for item in exp_items:
            title_date_line(doc, item["title"], item["org"], item["dates"], style)
            for b in item["bullets"]:
                bullet(doc, b)
    else:
        body_para(doc, "Fresher / Student — focused on academic learning and personal projects.", color=gray)

    project_items = parse_block_lines(data.get("projects", ""))
    if project_items:
        section_heading(doc, "Projects", style)
        for item in project_items:
            title_date_line(doc, item["title"], item["org"], item["dates"], style)
            for b in item["bullets"]:
                bullet(doc, b)

    edu_items = parse_block_lines(data.get("education", ""))
    if edu_items:
        section_heading(doc, "Education", style)
        for item in edu_items:
            title_date_line(doc, item["title"], item["org"], item["dates"], style)
            for b in item["bullets"]:
                body_para(doc, b, size=10, color=gray)

    achievements = [a.strip() for a in (data.get("achievements") or "").split("\n") if a.strip()]
    if achievements:
        section_heading(doc, "Achievements & Certifications", style)
        body_para(doc, "  ·  ".join(achievements))

    extra = []
    if data.get("availability"):
        extra.append(f"Availability: {data.get('availability')}")
    if data.get("interests"):
        extra.append(f"Interests: {data.get('interests')}")
    if data.get("references"):
        extra.append(f"References: {data.get('references')}")
    if extra:
        section_heading(doc, "Additional", style)
        body_para(doc, "  ".join(extra))

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio.read()